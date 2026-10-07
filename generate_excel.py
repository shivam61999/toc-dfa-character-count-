import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = [
    [1, "Project Creation & Stack Setup", 
     "Create a responsive web application for an Exact Character Count DFA Validator using React, TypeScript, Vite, and Tailwind CSS. Structure the project into clean components, types, automata utilities, and styles.", 
     "Complete working React + Vite project architecture", 
     "Main project architecture"],

    [2, "Formal 5-Tuple Definition", 
     "Formally specify the DFA as a 5-tuple M = (Q, Σ, δ, q₀, F) to recognize strings over Σ = {a, b} where the count of 'a's is 3m and count of 'b's is 2n for non-negative integers m, n.", 
     "Formal mathematical 5-tuple specification", 
     "Theoretical model"],

    [3, "Modular State Space Design", 
     "Construct the 6 states Q = {q₀,₀, q₁,₀, q₂,₀, q₀,₁, q₁,₁, q₂,₁} where each state q_{i,j} tracks the remainder residues i = Nₐ mod 3 and j = Nb mod 2.", 
     "6-state modular residue state space", 
     "Core automata model"],

    [4, "Cartesian Product Machine", 
     "Model the machine as a Cartesian product automaton M = M₁ × M₂ where M₁ computes Nₐ mod 3 with 3 states and M₂ computes Nb mod 2 with 2 states, proving |Q| = 3 × 2 = 6.", 
     "Product automaton M₁ × M₂ design", 
     "Mathematical theory"],

    [5, "Deterministic Transition Function", 
     "Implement the state transition function δ: Q × Σ → Q where δ(q_{i,j}, a) = q_{(i+1) mod 3, j} and δ(q_{i,j}, b) = q_{i, (j+1) mod 2}.", 
     "Accurate deterministic transition logic", 
     "Core algorithm"],

    [6, "Empty String Acceptance (ε)", 
     "Validate and handle the empty string ε such that for m=0 and n=0, Nₐ=0 (3×0) and Nb=0 (2×0) are valid non-negative multiples, designating q₀,₀ as both initial and accepting state.", 
     "ε acceptance logic at start state q₀,₀", 
     "Boundary verification"],

    [7, "Commutative Order Invariance", 
     "Ensure that the relative order and interleaving of 'a's and 'b's does not affect the halting state due to the abelian (commutative) nature of modular addition.", 
     "Order-invariant simulation results (aaabb, bbaaa, ababa)", 
     "Algorithm correctness"],

    [8, "Dead / Trap State Handling", 
     "Implement an explicit dead/trap state q_trap that captures any symbol σ ∉ {a, b}, loops on all subsequent symbols, and permanently rejects the string.", 
     "Trap state routing and illegal character error messages", 
     "Error handling & robustness"],

    [9, "Transition Table Visualization", 
     "Build an interactive 2D State Transition Table showing all states, mod remainders, next states on 'a' and 'b', and state classifications.", 
     "Formal state transition table (δ: Q × Σ → Q)", 
     "Visualization module"],

    [10, "Active Row and Cell Highlighting", 
     "Dynamically highlight the current state row and the active input column ('a' or 'b') in the transition table in real time as the simulation steps through input.", 
     "Live transition table feedback synchronized with simulation", 
     "Visualization"],

    [11, "Interactive SVG Automaton Graph", 
     "Create a responsive SVG state machine diagram with a 2-row × 3-column layout, start arrow pointing to q₀,₀, double concentric circle on accept state q₀,₀, and curved transition edges.", 
     "Vector state machine diagram with proper topology", 
     "Diagram visualization"],

    [12, "Active State Glow & Halos", 
     "Animate the currently active state in the SVG diagram with glowing borders, radial ping animations, and state labels to track runtime execution.", 
     "Live active state visualization with visual feedback", 
     "Animation & UI"],

    [13, "Active Edge Traveling Pulse", 
     "Highlight the traversed transition edge and animate an active traveling pulse along the directed curve corresponding to the read character 'a' or 'b'.", 
     "Animated active transition paths", 
     "Visualization"],

    [14, "State Inspection Modal", 
     "Allow users to click any state node in the diagram to open a modal detailing its mathematical invariant, residues, and transition behavior.", 
     "State inspector modal with formal invariants", 
     "Educational UI"],

    [15, "Input Tape Visualizer", 
     "Build an input tape displaying characters in individual cells, distinguishing processed, current, and upcoming characters with color coding.", 
     "Interactive tape display with distinct cell states", 
     "Visualization"],

    [16, "Dynamic Read Head", 
     "Implement a visual tape pointer head (▼) that tracks the exact character index currently being evaluated by the DFA.", 
     "Dynamic read head pointer", 
     "Simulation UI"],

    [17, "Live Modulo Invariant Tracker", 
     "Display real-time counters and integer equations for Nₐ = 3m + rₐ and Nb = 2n + rb alongside multipliers m and n.", 
     "Live invariant statistics panel with mathematical formulas", 
     "Analytical tracking"],

    [18, "Step-by-Step Simulation Engine", 
     "Implement simulation playback controls including Play, Pause, Step Next, Step Backward, Jump to Start, Jump to End, and Reset.", 
     "Controllable execution engine for step-by-step trace", 
     "Simulation engine"],

    [19, "Simulation Speed & Shortcuts", 
     "Provide playback speed toggles (0.5x to 3x) and keyboard hotkeys (Space to play/pause, Arrow keys to step, R to reset).", 
     "Speed controls and keyboard navigation", 
     "User experience"],

    [20, "Instantaneous Descriptions (IDs)", 
     "Generate formal step-by-step Instantaneous Descriptions (q, w) ⊢ (q', w') representing the configuration sequence, with one-click copy to clipboard.", 
     "Copyable formal mathematical configuration trace", 
     "Academic reporting"],

    [21, "Comprehensive Test Suite", 
     "Preload a categorized batch test suite of 24+ test cases covering ε, minimal accepted strings, higher multipliers, defective 'a's, defective 'b's, and trap characters.", 
     "Categorized test suite across all boundary conditions", 
     "Verification"],

    [22, "Automated Pass/Fail Audit", 
     "Run all test cases automatically against their theoretical expectations and display a real-time pass rate percentage and audit summary.", 
     "100% pass-rate scorecard with pass/fail badges", 
     "Testing & QA"],

    [23, "Custom Test Case Addition", 
     "Allow users to input custom strings and descriptions to dynamically add them to the test suite and evaluate their acceptance.", 
     "Custom test creator with instant validation", 
     "Testing"],

    [24, "Random Test String Generator", 
     "Build a generator to produce randomized valid strings with custom (m, n) multipliers or inject specific modular defects for testing.", 
     "Random test generator modal with Fisher-Yates interleaving", 
     "Test generation"],

    [25, "Verdict Banner & Confetti", 
     "Display a prominent status banner showing ACCEPTED or REJECTED with full algebraic reasoning and celebratory confetti on string acceptance.", 
     "Visual verdict banner & celebration animation", 
     "User experience"],

    [26, "Myhill-Nerode Minimality Proof", 
     "Include a formal mathematical proof showing that all 6 states belong to pairwise distinguishable equivalence classes, proving the DFA is strictly minimal.", 
     "State minimality proof section in academic dossier", 
     "Academic dossier"],

    [27, "Viva Voce Q&A Dossier", 
     "Compile an interactive accordion of common oral examination and viva voce questions covering language index, commutativity, NFAs, and trap states.", 
     "Interactive Viva Q&A section with toggleable answers", 
     "Viva preparation"],

    [28, "Real-World Applications Module", 
     "Document real-world engineering applications of modulo DFAs in bioinformatics codon reading frames, hardware parity checking, VLSI clock dividers, and compiler memory alignment.", 
     "Real-world applications cards with academic references", 
     "Academic dossier"],

    [29, "Progressive Web App (PWA) Offline", 
     "Configure vite-plugin-pwa with Service Worker caching and web manifest to enable 100% offline functionality and native installation on Android and Desktop.", 
     "Offline PWA with install modal and cache precaching", 
     "Cross-platform deployment"],

    [30, "Multi-Platform Deployment Workflows", 
     "Configure project with relative asset paths, vercel.json, and GitHub Actions workflow for seamless 1-click deployment to GitHub Pages, Vercel, and Render.", 
     "Universal deployment configurations and guides", 
     "DevOps & Hosting"]
]

headers = ["No.", "Category", "Prompt", "Expected Output", "Use"]

# 1. Write CSV file
csv_filename = "TOC_DFA_Prompt_Engineering_Lifecycle.csv"
with open(csv_filename, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(headers)
    writer.writerows(data)
print(f"CSV generated: {csv_filename}")

# 2. Write Excel XLSX file with professional styling
xlsx_filename = "TOC_DFA_Prompt_Engineering_Lifecycle.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "DFA Prompt Lifecycle"

# Header formatting
header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Dark Indigo
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Data formatting
regular_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
center_alignment = Alignment(horizontal="center", vertical="center")
left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# Write Header
ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border
ws.row_dimensions[1].height = 28

# Write Data Rows
for row_idx, row_data in enumerate(data, start=2):
    ws.append(row_data)
    fill = zebra_fill if row_idx % 2 == 0 else white_fill
    
    # Cell 1: No.
    c1 = ws.cell(row=row_idx, column=1)
    c1.alignment = center_alignment
    c1.font = bold_font
    c1.fill = fill
    c1.border = thin_border
    
    # Cell 2: Category
    c2 = ws.cell(row=row_idx, column=2)
    c2.alignment = left_alignment
    c2.font = bold_font
    c2.fill = fill
    c2.border = thin_border
    
    # Cell 3: Prompt
    c3 = ws.cell(row=row_idx, column=3)
    c3.alignment = left_alignment
    c3.font = regular_font
    c3.fill = fill
    c3.border = thin_border
    
    # Cell 4: Expected Output
    c4 = ws.cell(row=row_idx, column=4)
    c4.alignment = left_alignment
    c4.font = regular_font
    c4.fill = fill
    c4.border = thin_border
    
    # Cell 5: Use
    c5 = ws.cell(row=row_idx, column=5)
    c5.alignment = left_alignment
    c5.font = regular_font
    c5.fill = fill
    c5.border = thin_border
    
    ws.row_dimensions[row_idx].height = 42

# Column widths
col_widths = {
    1: 8,   # No.
    2: 28,  # Category
    3: 56,  # Prompt
    4: 34,  # Expected Output
    5: 24   # Use
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# Enable AutoFilter
ws.auto_filter.ref = f"A1:E{len(data)+1}"

wb.save(xlsx_filename)
print(f"XLSX generated: {xlsx_filename}")
